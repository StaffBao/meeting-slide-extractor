import py_compile
import os
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


class PackageSmokeTests(unittest.TestCase):
    def test_required_distribution_files_exist(self):
        required = [
            "SKILL.md",
            "README.md",
            "SOURCE_VIDEO_NOTICE.md",
            "DISCLAIMER.md",
            "LICENSE",
            "requirements.txt",
            "agents/openai.yaml",
            "references/env.md",
        ]
        for relative_path in required:
            with self.subTest(path=relative_path):
                self.assertTrue((ROOT / relative_path).is_file())

    def test_all_scripts_compile(self):
        scripts = sorted((ROOT / "scripts").glob("*.py"))
        self.assertGreater(len(scripts), 0)
        for script in scripts:
            with self.subTest(script=script.name):
                py_compile.compile(str(script), doraise=True)

    def test_word_transcript_contains_overview_and_sections(self):
        sys.path.insert(0, str(ROOT / "scripts"))
        try:
            from build_notes_ppt import build_transcript_docx
            from docx import Document
            from docx.oxml.ns import qn
            from docx.shared import Pt

            meta = [
                {"n": 1, "t": 0, "notes": "这是第一段发言。"},
                {"n": 2, "t": 30, "notes": "这是第二段发言。\n\n这是新段落。"},
            ]
            sections = [{
                "idx": 1,
                "speaker": "张三",
                "topic": "项目进展",
                "title": "张三 项目进展",
                "from_n": 1,
                "to_n": 2,
                "from_t": 0,
                "to_t": 60,
            }]
            tmp_parent = Path(os.environ.get("MEETING_TEST_TMP", ROOT))
            tmp_parent.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryDirectory(dir=tmp_parent) as tmp:
                output = Path(tmp) / "transcript_by_speaker.docx"
                build_transcript_docx(meta, str(output), sections=sections)
                self.assertTrue(output.is_file())
                doc = Document(output)
                text = "\n".join(p.text for p in doc.paragraphs)
                self.assertIn("逐字稿", text)
                self.assertIn("第 1 节 张三 项目进展", text)
                self.assertIn("这是第一段发言。", text)
                self.assertEqual(len(doc.tables), 1)
                self.assertEqual(doc.tables[0].cell(1, 1).text, "张三")
                normal = doc.styles["Normal"]
                self.assertEqual(normal.font.name, "Arial Unicode MS")
                self.assertEqual(normal.font.size, Pt(12))
                self.assertEqual(
                    normal._element.rPr.rFonts.get(qn("w:eastAsia")), "宋体")
        finally:
            sys.path.remove(str(ROOT / "scripts"))


if __name__ == "__main__":
    unittest.main()
