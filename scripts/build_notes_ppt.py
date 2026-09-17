# build_notes_ppt.py  (V2.2.0)
# Shared library for building deliverables from slides_meta.json (which must
# carry, per slide: n, file, t, and an optional notes string).
# Used by align_by_time.py and punctuate_notes.py.
#
# Standard deliverables:
#   - output/slides_viewer.html        (V2.2.0: sectioned, left thumbnail rail,
#     page-jump box, in-page editor, section titles editable — ONE
#     self-contained .html by default)
#   - output/slides.pptx               (clean deck, image only)
#   - output/transcript_by_speaker.docx (formatted Word transcript grouped by speaker)
# Diagnostic/compatibility outputs:
#   - output/transcript_from_slides.md (per-slide OCR + notes, with section
#     headings so the transcript carries the same 【节】 markers as the deck)
#   - output/slides_with_notes.pptx    (legacy opt-in only via --notes-pptx)
import argparse, json, os, re, html, hashlib, base64
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Part of the draft-storage salt. Bumping it on every release guarantees that a
# draft written by an older build of the same deck can never be restored into a
# newer one (same title + page count + file name would otherwise collide).
DECK_VERSION = "2.2.0"

ACCENT = RGBColor(0x25, 0x63, 0xEB)
INK = RGBColor(0x15, 0x22, 0x38)
MUTED = RGBColor(0x6B, 0x76, 0x87)


def load_meta(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def save_meta(path, meta):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)


def parse_ocr(path):
    """Parse output/slides_ocr.txt -> {n: text}. Section header: '# 第 N 页'."""
    if not path or not os.path.exists(path):
        return {}
    out = {}
    current = None
    buf = []
    pat = re.compile(r"# 第\s*(\d+)\s*页")
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.rstrip("\n")
            m = pat.match(line)
            if m:
                if current is not None:
                    out[current] = "\n".join(buf).strip()
                current = int(m.group(1))
                buf = []
            else:
                buf.append(line)
    if current is not None:
        out[current] = "\n".join(buf).strip()
    return out


# --------------------------------------------------------------- sections ---
def load_sections(obj):
    """Accept a path to sections.json, an already-loaded dict, or None.

    Returns a normalised list of section dicts sorted by first slide, or [].
    """
    if not obj:
        return []
    if isinstance(obj, str):
        if not os.path.exists(obj):
            return []
        with open(obj, encoding="utf-8") as f:
            obj = json.load(f)
    secs = obj.get("sections") if isinstance(obj, dict) else obj
    if not secs:
        return []
    out = []
    for i, s in enumerate(secs, 1):
        out.append({
            "idx": int(s.get("idx", i)),
            "kind": s.get("kind", "talk"),
            "speaker": s.get("speaker") or "",
            "topic": s.get("topic") or "",
            "title": s.get("title") or s.get("speaker") or ("节 %d" % i),
            "from_n": int(s.get("from_n", 1)),
            "to_n": int(s.get("to_n", s.get("from_n", 1))),
            "from_t": s.get("from_t"),
            "to_t": s.get("to_t"),
            "anchor": s.get("anchor") or "",
            "anchor_t": s.get("anchor_t"),
        })
    out.sort(key=lambda s: s["from_n"])
    return out


def load_name_fixes(obj):
    """ASR-spelling -> slide-spelling map published by build_sections.py.

    The slide is authoritative for a speaker's name (measured: 史兰 vs 史岚,
    蒲一虎 vs 浦义虎). Applying the same substitution to the notes keeps the
    deck, the viewer and the transcript saying the same thing.
    """
    if not obj:
        return {}
    if isinstance(obj, str):
        if not os.path.exists(obj):
            return {}
        with open(obj, encoding="utf-8") as f:
            obj = json.load(f)
    if isinstance(obj, dict):
        return {k: v for k, v in (obj.get("name_fixes") or {}).items() if k and v and k != v}
    return {}


def apply_name_fixes(meta, fixes):
    """Replace mis-heard speaker names in every slide's notes + OCR-free text.

    Returns the number of replacements actually made.
    """
    if not fixes:
        return 0
    pat = re.compile("|".join(re.escape(k) for k in sorted(fixes, key=len, reverse=True)))
    hits = 0

    def sub(m):
        nonlocal hits
        hits += 1
        return fixes[m.group(0)]

    for s in meta:
        if s.get("notes"):
            s["notes"] = pat.sub(sub, s["notes"])
    return hits


def _norm_cjk(s):
    return re.sub(r"[\s\u3000]+", "", s or "")


def _lcs_span(a, b):
    """Longest common substring -> (length, i_end_in_a, j_end_in_b), or None.

    The notes are LLM-punctuated while the anchor comes from the raw ASR, so the
    hand-over sentence never matches verbatim; a longest-common-substring scan is
    what survives the inserted punctuation. Inputs are short (<400 chars).
    """
    if not a or not b:
        return None
    prev = [0] * (len(b) + 1)
    best = (0, 0, 0)
    for i in range(1, len(a) + 1):
        cur = [0] * (len(b) + 1)
        ai = a[i - 1]
        for j in range(1, len(b) + 1):
            if ai == b[j - 1]:
                n = prev[j - 1] + 1
                cur[j] = n
                if n > best[0]:
                    best = (n, i, j)
        prev = cur
    return best if best[0] else None


# sentence ends used to snap the divider onto a real boundary
SENT_END = "。！？；!?;\n"


def _apply_fixes(text, fixes):
    if not fixes:
        return text
    for k, v in fixes.items():
        text = text.replace(k, v)
    return text


def _split_after(raw, anchor, min_match=8, fixes=None):
    """Char offset just AFTER the anchor's sentence inside `raw`, or None.

    The divider belongs after the host's hand-over sentence: the bridge stays
    with the section that is ending, and the new section opens cleanly with the
    next speaker's content.

    Both sides are passed through `fixes` first, because the notes will already
    carry the corrected speaker name while the anchor (raw ASR) still has the
    mis-heard one — otherwise the substitution would shorten the common run
    below the threshold.
    """
    span = _lcs_span(_norm_cjk(_apply_fixes(anchor or "", fixes)),
                     _norm_cjk(_apply_fixes(raw, fixes)))
    if not span or span[0] < min_match:
        return None
    n, _, j_end = span
    # j_end indexes the space-stripped copy; map it back onto the raw string by
    # walking and counting non-space characters.
    seen, raw_end = 0, len(raw)
    for i, ch in enumerate(raw):
        if ch.isspace() or ch == "\u3000":
            continue
        seen += 1
        if seen == j_end:
            raw_end = i + 1
            break
    k = raw_end
    while k < len(raw) and raw[k] not in SENT_END:
        k += 1
    if k < len(raw) and raw[k] in "。！？!?":
        k += 1
    # do not strand a lone closing quote / bracket on the far side of the split
    while k < len(raw) and raw[k] in " \t\u3000\u201d\u2019\"'）)】》”’":
        k += 1
    return k if k > 0 else None


def locate_anchors(meta, sections, min_match=8, fixes=None):
    """Find where each section's hand-over sits inside the PREVIOUS page's notes,
    so the divider can be placed mid-page instead of only at a slide boundary.

    Returns {section_idx: {"page": n, "off": char_offset}} where `off` is the
    character offset in that page's raw notes after which the marker goes.
    Sections whose anchor cannot be located simply get a card-level divider.
    """
    by_n = {s["n"]: s for s in meta}
    out = {}
    for sec in sections:
        anchor = sec.get("anchor")
        if not anchor:
            continue
        t = sec.get("anchor_t")
        start = None
        if t is not None:
            for m in meta:
                if m.get("t", 0) <= t:
                    start = m["n"]
                else:
                    break
        if start is None:
            start = sec["from_n"]
        if start > sec["from_n"]:
            start = sec["from_n"]
        # earliest page wins: the hand-over is spoken BEFORE the speaker advances
        for n in range(start, sec["from_n"] + 1):
            s = by_n.get(n)
            if not s:
                continue
            raw = _norm_notes(s.get("notes", ""))
            off = _split_after(raw, anchor, min_match, fixes)
            if off:
                out[sec["idx"]] = {"page": n, "off": off}
                break
    return out


def fmt_mmss(sec):
    """mm:ss, or h:mm:ss once the recording passes an hour.

    Training recordings routinely run for hours, and formatting 2h38m as
    "158:34" reads as a broken minute field.
    """
    if sec is None:
        return ""
    try:
        sec = float(sec)
    except Exception:
        return ""
    if sec < 0:
        sec = 0
    h = int(sec // 3600)
    m = int((sec % 3600) // 60)
    s = int(sec % 60)
    if h:
        return "%d:%02d:%02d" % (h, m, s)
    return "%02d:%02d" % (m, s)


def section_range_label(meta, sec):
    a = fmt_mmss(sec.get("from_t"))
    to_t = sec.get("to_t")
    b = fmt_mmss(to_t) if to_t else ""      # 0/None -> section end unknown
    span = "第 %d–%d 页" % (sec["from_n"], sec["to_n"])
    if a and b:
        span += " · %s–%s" % (a, b)
    elif a:
        span += " · %s 起" % a
    return span


# ------------------------------------------------------------- pptx output ---
def _fit_box(iw, ih, bw, bh):
    ar = (iw / ih) if ih else 16 / 9
    box_ar = bw / bh
    if ar > box_ar:
        w, h = bw, int(bw / ar)
    else:
        h, w = bh, int(bh * ar)
    return w, h, int((bw - w) / 2), int((bh - h) / 2)


def _img_size(path, default=(16, 9)):
    from PIL import Image
    try:
        return Image.open(path).size
    except Exception:
        return default


def _add_section_slide(prs, meta, sec):
    """A 【新增节】 divider slide: accent band, section title, range + speaker."""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    SW, SH = prs.slide_width, prs.slide_height
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, int(SW * 0.011), SH)
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()
    band.shadow.inherit = False
    box = slide.shapes.add_textbox(int(SW * 0.085), int(SH * 0.28), int(SW * 0.83), int(SH * 0.46))
    tf = box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "新增节"
    r0.font.size = Pt(16)
    r0.font.bold = True
    r0.font.color.rgb = ACCENT
    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = sec["title"]
    r1.font.size = Pt(30)
    r1.font.bold = True
    r1.font.color.rgb = INK
    p1.space_before = Pt(10)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = section_range_label(meta, sec)
    r2.font.size = Pt(13)
    r2.font.color.rgb = MUTED
    p2.space_before = Pt(14)
    # No notes on the divider either: the PPT deliverable is "slides only", so the
    # deck must stay free of any transcript-adjacent text.
    return slide


def _add_image_slide(prs, s, with_notes=True):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    SW, SH = prs.slide_width, prs.slide_height
    iw, ih = _img_size(s["file"])
    w, h, x, y = _fit_box(iw, ih, SW, SH)
    slide.shapes.add_picture(s["file"], x, y, width=w, height=h)
    if with_notes:
        notes = s.get("notes", "")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    return slide


def build_notes_pptx(meta, out_path, sections=None):
    """Deck of slide images + speaker notes. When `sections` is given, a
    【新增节】 divider slide is inserted in front of each section's first slide
    and PowerPoint's own section pane is populated as well."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    secs = load_sections(sections)
    by_from = {s["from_n"]: s for s in secs}
    for s in meta:
        sec = by_from.get(s["n"])
        if sec:
            _add_section_slide(prs, meta, sec)
        _add_image_slide(prs, s, with_notes=True)
    prs.save(out_path)
    if secs:
        _apply_native_sections(out_path, prs, secs)
    return out_path


def _apply_native_sections(out_path, prs, secs):
    """Write PowerPoint's native `p14:sectionLst` so the 节 also show up in
    PowerPoint's section pane (幻灯片 → 节).

    python-pptx has no API for this, so the extension XML is written by hand.
    It is strictly non-destructive: the result is staged to a temp file, re-opened
    as a structural sanity check, and only then moved over the target. Anything
    going wrong leaves the already-saved deck (with its visible 【新增节】 divider
    slides) exactly as it was.
    """
    import uuid
    from lxml import etree
    from pptx import Presentation as _P

    NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    EXT_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
    tmp = out_path + ".sect.tmp"
    # 显式注册前缀，否则 lxml 会给这个命名空间自造一个 "ns0:" 前缀——XML 合法，
    # 但读起来不像规范的 p14:，排查时容易误判成「没写进去」。
    etree.register_namespace("p14", NS_P14)

    def q(ns, tag):
        return "{%s}%s" % (ns, tag)

    try:
        pres = prs.part._element
        sld_ids = pres.findall(q(NS_P, "sldIdLst") + "/" + q(NS_P, "sldId"))
        ids = [e.get("id") for e in sld_ids]
        # Slide order in the pptx == the order we added them: for every section
        # one divider slide followed by that section's k image slides.
        groups, cursor = [], 0
        for sec in secs:
            k = sec["to_n"] - sec["from_n"] + 1
            chunk = ids[cursor:cursor + 1 + k]
            if len(chunk) != 1 + k:
                return False                        # layout mismatch -> keep simple
            groups.append({"sec": sec, "ids": chunk})
            cursor += 1 + k
        if cursor != len(ids):
            return False

        ext_lst = pres.find(q(NS_P, "extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(pres, q(NS_P, "extLst"))
        ext = etree.SubElement(ext_lst, q(NS_P, "ext"))
        ext.set("uri", EXT_URI)
        lst = etree.SubElement(ext, q(NS_P14, "sectionLst"))
        for g in groups:
            se = etree.SubElement(lst, q(NS_P14, "section"))
            se.set("name", g["sec"]["title"])
            se.set("id", "{%s}" % str(uuid.uuid4()).upper())
            sl = etree.SubElement(se, q(NS_P14, "sldIdLst"))
            for sid in g["ids"]:
                e = etree.SubElement(sl, q(NS_P14, "sldId"))
                e.set("id", str(sid))

        prs.save(tmp)
        _P(tmp)                                     # cheap structural sanity check
        os.replace(tmp, out_path)
        return True
    except Exception as exc:              # noqa: BLE001 - best effort by design
        print("  ! 原生节（PowerPoint 节窗格）写入失败，已保留可见的【新增节】分页：%s" % exc)
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except Exception:
            pass
        return False


def build_clean_pptx(meta, out_path, sections=None):
    """THE PPT deliverable: slides only, contain-fit, no notes / no side text.

    V2.2.0: pass `sections` to insert a 【新增节】 divider slide in front of each
    section's first slide (and to populate PowerPoint's own 节 pane). The deck
    still carries NO transcript — the transcript lives in the HTML viewer, which
    is the point of the V2.x split: 一份纯幻灯片 + 一份幻灯片与逐字稿镶嵌的页面。
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    secs = load_sections(sections)
    by_from = {s["from_n"]: s for s in secs}
    for s in meta:
        sec = by_from.get(s["n"])
        if sec:
            _add_section_slide(prs, meta, sec)
        _add_image_slide(prs, s, with_notes=False)
    prs.save(out_path)
    if secs:
        _apply_native_sections(out_path, prs, secs)
    return out_path


# ------------------------------------------------------------- markdown ------
def build_transcript(meta, out_path, sections=None, title=None, with_time=True):
    """旧版 Markdown 逐字稿生成器，保留供兼容调用。

    刻意不含幻灯片 OCR：OCR 是「把 PPT 转成文字」，那是另一份东西。混进逐字稿会让人
    分不清哪句是讲出来的、哪句是从画面扫出来的；PPT 的文字本身就在 PPTX 和 HTML 里。
    也不做逐页切分——逐字稿的阅读单位是「谁在讲」，页边界属于 viewer 的职责。

    节内的正文按页顺序拼接（讲稿是按时长对齐到页的），页间的空行由
    `_notes_paragraphs` 统一规整，因此拼接后仍是自然段落而不是碎片。
    """
    secs = load_sections(sections)
    lines = ["# %s\n\n" % (title or "逐字稿（按讲者分节）")]
    by_n = {s["n"]: s for s in meta}

    def emit(n0, n1):
        for n in range(n0, n1 + 1):
            s = by_n.get(n)
            if not s:
                continue
            for p in _notes_paragraphs(s.get("notes", "")):
                lines.append(p + "\n\n")

    if not secs:
        emit(min(by_n) if by_n else 1, max(by_n) if by_n else 0)
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("".join(lines))
        return out_path

    lines.append("## 讲者一览\n\n")
    lines.append("| 节 | 讲者 | 题目 | 范围 |\n")
    lines.append("|---|---|---|---|\n")
    for s in secs:
        lines.append("| %d | %s | %s | %s |\n" % (
            s["idx"], s.get("speaker") or "（未识别）", s.get("topic") or "—",
            section_range_label(meta, s)))
    lines.append("\n---\n\n")
    for sec in secs:
        lines.append("## 节 %d · %s\n\n" % (sec["idx"], sec["title"]))
        if with_time:
            lines.append("*%s*\n\n" % section_range_label(meta, sec))
        emit(sec["from_n"], sec["to_n"])
        lines.append("\n---\n\n")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("".join(lines))
    return out_path


def _docx_set_font(target, latin="Arial Unicode MS", east_asia="宋体", size=None,
                   bold=None, color=None):
    """Apply fonts consistently to either a style or a run."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    target.font.name = latin
    if size is not None:
        target.font.size = Pt(size)
    if bold is not None:
        target.font.bold = bold
    if color is not None:
        target.font.color.rgb = color
    rpr = target._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _docx_set_cell_shading(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _docx_set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start),
                        ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _docx_set_table_borders(table, color="D9D9D9", size="6"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def _docx_repeat_header(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def build_transcript_docx(meta, out_path, sections=None, title=None, with_time=True):
    """Build the standard Word transcript deliverable.

    The document contains speech only, grouped continuously by speaker/section.
    Slide OCR remains in the diagnostic Markdown output and is never mixed into
    this transcript.
    """
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor as DocxRGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = document.styles
    # Transcript typography: Chinese uses SimSun (宋体), Latin text uses
    # Arial Unicode MS. Body text is fixed at 12 pt.
    _docx_set_font(styles["Normal"], size=12, color=DocxRGBColor(0, 0, 0))
    styles["Normal"].paragraph_format.line_spacing = 1.45
    styles["Normal"].paragraph_format.space_after = Pt(6)
    _docx_set_font(styles["Title"], size=22, bold=True,
                   color=DocxRGBColor(0, 0, 0))
    styles["Title"].paragraph_format.space_after = Pt(14)
    # Some Word templates attach a blue bottom border to the built-in Title
    # style. The transcript uses whitespace for separation instead.
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    _docx_set_font(styles["Heading 1"], size=16, bold=True,
                   color=DocxRGBColor(0, 0, 0))
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    _docx_set_font(styles["Heading 2"], size=13, bold=True,
                   color=DocxRGBColor(0, 0, 0))
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)

    document.core_properties.title = title or "逐字稿 按讲者分节"
    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.add_run(title or "逐字稿 按讲者分节")

    secs = load_sections(sections)
    by_n = {s["n"]: s for s in meta}

    def emit(n0, n1):
        emitted = 0
        for n in range(n0, n1 + 1):
            slide = by_n.get(n)
            if not slide:
                continue
            for text in _notes_paragraphs(slide.get("notes", "")):
                paragraph = document.add_paragraph(style="Normal")
                paragraph.paragraph_format.widow_control = True
                paragraph.add_run(text)
                emitted += 1
        if not emitted:
            paragraph = document.add_paragraph(style="Normal")
            run = paragraph.add_run("（未识别到发言正文）")
            run.italic = True
            run.font.color.rgb = DocxRGBColor(100, 100, 100)

    if not secs:
        emit(min(by_n) if by_n else 1, max(by_n) if by_n else 0)
        document.save(out_path)
        return out_path

    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("讲者一览")
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(0.55), Inches(1.25), Inches(3.25), Inches(1.65))
    headers = ("节", "讲者", "题目", "范围")
    for idx, (cell, label, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _docx_set_cell_shading(cell, "1F4E78")
        _docx_set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        _docx_set_font(run, size=12, bold=True,
                       color=DocxRGBColor(255, 255, 255))
    _docx_repeat_header(table.rows[0])

    for row_index, sec in enumerate(secs, start=1):
        row = table.add_row()
        values = (
            str(sec["idx"]),
            sec.get("speaker") or "（未识别）",
            sec.get("topic") or "—",
            section_range_label(meta, sec),
        )
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _docx_set_cell_margins(cell)
            if row_index % 2 == 0:
                _docx_set_cell_shading(cell, "F3F7FB")
            p = cell.paragraphs[0]
            p.alignment = (WD_ALIGN_PARAGRAPH.CENTER
                           if col_index in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            _docx_set_font(run, size=12, color=DocxRGBColor(0, 0, 0))
    _docx_set_table_borders(table)

    for sec in secs:
        document.add_page_break()
        heading = document.add_paragraph(style="Heading 1")
        heading.add_run("第 %d 节 %s" % (sec["idx"], sec["title"]))
        if with_time:
            range_p = document.add_paragraph()
            range_p.paragraph_format.space_after = Pt(10)
            range_run = range_p.add_run(section_range_label(meta, sec))
            range_run.italic = True
            _docx_set_font(range_run, size=12, color=DocxRGBColor(89, 89, 89))
        emit(sec["from_n"], sec["to_n"])

    document.save(out_path)
    return out_path


def build_markdown(meta, ocr_by_slide, out_path, sections=None, title=None):
    """流水线的**自查/对照**产物：每页 OCR + 对应讲稿。

    注意这与交付物 `build_transcript` 不同——这里刻意把幻灯片 OCR 和讲稿并排放，
    是为了人工复核「画面上的文字」与「说出来的话」是否对得上。
    """
    secs = load_sections(sections)
    by_from = {s["from_n"]: s for s in secs}
    lines = ["# %s\n" % (title or "幻灯片 OCR 与对应讲稿（按录像时间排列）"), "\n"]
    if secs:
        lines.append("## 节次一览\n\n")
        for s in secs:
            lines.append("- **节 %d** %s · %s\n" % (s["idx"], s["title"],
                                                 section_range_label(meta, s)))
        lines.append("\n")
    for s in meta:
        n = s["n"]
        if n in by_from:
            sec = by_from[n]
            lines.append("\n---\n\n")
            lines.append("# 【新增节 %d】%s\n\n" % (sec["idx"], sec["title"]))
            lines.append("> %s\n\n" % section_range_label(meta, sec))
        t = s.get("t", 0)
        lines.append("## 第 %03d 页 · 录像 %s\n\n" % (n, fmt_mmss(t) or t))
        lines.append("### 幻灯片 OCR\n\n")
        lines.append(ocr_by_slide.get(n, "（无识别结果）") + "\n\n")
        lines.append("### 对应讲稿\n\n")
        notes = s.get("notes", "")
        lines.append((notes if notes else "（未匹配到讲稿）") + "\n\n")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("".join(lines))
    return out_path


# ------------------------------------------------------------------ shared ---
def _norm_notes(text):
    """Normalise a raw notes string: CRLF/CR -> LF, trimmed.

    Used both for rendering paragraphs and for the data-notes attribute that
    the in-page editor reads, so the two can never drift apart.
    """
    if not text:
        return ""
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def _notes_paragraphs(text):
    """Split a notes string into paragraph strings (blank line = paragraph)."""
    text = _norm_notes(text)
    if not text:
        return []
    out = []
    for block in re.split(r"\n\s*\n", text):
        block = " ".join(x.strip() for x in block.split("\n") if x.strip())
        if block:
            out.append(block)
    return out


_HTML_TMPL = r"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title>
<style>
:root{--bg:#eef1f5;--card:#fff;--line:#dfe4ea;--ink:#1b1f26;--muted:#6b7687;
      --accent:#2563eb;--stage:#12161d;--head:#eef2f7;--track:#e7ebf0;--thumb:#b9c3d0;
      --ok:#16a34a;--warn:#b45309;--amber:#f0b429;}
*{box-sizing:border-box}
html{scroll-behavior:smooth}
body{margin:0;background:var(--bg);color:var(--ink);
     font-family:"Microsoft YaHei",system-ui,-apple-system,"Segoe UI",sans-serif}
header.top{position:sticky;top:0;z-index:50;background:rgba(255,255,255,.95);
     backdrop-filter:blur(8px);border-bottom:1px solid var(--line);
     display:flex;align-items:center;gap:9px;padding:9px 16px;flex-wrap:wrap}
header.top .t{font-weight:700;font-size:15px;white-space:nowrap;overflow:hidden;
     text-overflow:ellipsis;max-width:34vw}
header.top .sp{margin-left:auto}
header.top button{border:1px solid var(--line);background:#fff;border-radius:8px;
     padding:6px 11px;cursor:pointer;font-size:13px;color:var(--ink);white-space:nowrap}
header.top button:hover{background:var(--head)}
header.top button.pri{background:var(--accent);border-color:var(--accent);color:#fff}
header.top button.pri:hover{background:#1d4fd7}
header.top .c{color:var(--muted);font-size:13px;font-variant-numeric:tabular-nums;
     min-width:62px;text-align:right}
.jump{display:inline-flex;align-items:center;gap:5px}
.jump input{width:64px;border:1px solid var(--line);border-radius:8px;padding:6px 7px;
     font-size:13px;color:var(--ink);text-align:center;font-variant-numeric:tabular-nums;
     font-family:inherit}
.jump input:focus{outline:2px solid rgba(37,99,235,.25);border-color:var(--accent)}
.jump .of{font-size:12px;color:var(--muted);font-variant-numeric:tabular-nums}
.tools{display:none;align-items:center;gap:8px;width:100%;margin-top:4px;padding-top:9px;
     border-top:1px dashed var(--line);flex-wrap:wrap}
body.editing .tools{display:flex}
.tools .stat{font-size:12px;color:var(--muted);font-variant-numeric:tabular-nums}
.tools .badge{font-size:12px;padding:3px 9px;border-radius:999px;border:1px solid var(--line);color:var(--muted)}
.tools .badge.ok{color:var(--ok);border-color:#bbe7c9;background:#f0fdf4}
.tools .badge.warn{color:var(--warn);border-color:#f3ddb4;background:#fffbeb}
.restore{display:none;align-items:center;gap:10px;max-width:1560px;margin:14px auto 0;
     padding:10px 14px;border:1px solid #f3ddb4;background:#fffbeb;border-radius:10px;
     font-size:13px;color:#8a5a09}
.restore.on{display:flex}
.restore button{border:1px solid #f3ddb4;background:#fff;border-radius:8px;
     padding:5px 11px;cursor:pointer;font-size:13px;color:#8a5a09;margin-left:auto}
/* ---- layout: left thumbnail rail + reading column ------------------------- */
.layout{display:flex;align-items:flex-start;gap:16px;max-width:1560px;
     margin:0 auto;padding:20px 18px 70px}
.wrap{flex:1 1 auto;min-width:0;margin:0;padding:0}
.rail{flex:0 0 186px;position:sticky;top:62px;max-height:calc(100vh - 82px);
     overflow-y:auto;background:var(--card);border:1px solid var(--line);
     border-radius:12px;padding:8px 7px 12px;
     scrollbar-width:thin;scrollbar-color:var(--thumb) var(--track)}
.rail::-webkit-scrollbar{width:9px}
.rail::-webkit-scrollbar-track{background:var(--track);border-radius:8px}
.rail::-webkit-scrollbar-thumb{background:var(--thumb);border-radius:8px;border:2px solid var(--track)}
body.rail-hidden .rail{display:none}
.rail-sect{font-size:11px;font-weight:700;color:var(--accent);background:#f3f7ff;
     border:1px solid #dbe6fb;border-radius:7px;padding:5px 7px;margin:7px 2px 8px;
     line-height:1.5;cursor:pointer}
.rail-sect:hover{background:#e6eefe}
.rail-sect:first-child{margin-top:2px}
.rail-sect .n{display:inline-block;font-size:10px;color:#fff;background:var(--accent);
     border-radius:5px;padding:0 5px;margin-right:5px;font-weight:700}
.thumb{position:relative;display:block;border:2px solid transparent;border-radius:8px;
     overflow:hidden;margin:0 0 6px;cursor:pointer;background:var(--stage);line-height:0}
.thumb img{width:100%;display:block}
.thumb .tn{position:absolute;left:3px;bottom:3px;background:rgba(17,24,39,.74);color:#fff;
     font-size:10px;line-height:14px;border-radius:4px;padding:0 4px;font-variant-numeric:tabular-nums}
.thumb:hover{border-color:var(--thumb)}
.thumb.active{border-color:var(--accent);box-shadow:0 0 0 2px rgba(37,99,235,.2)}
/* ---- section divider between cards --------------------------------------- */
.sect{display:flex;align-items:center;gap:12px;margin:2px 0 22px;padding:13px 17px;
     border:1px solid #cfe0ff;border-left:5px solid var(--accent);border-radius:12px;
     background:linear-gradient(90deg,#f2f7ff,#fff 62%);
     box-shadow:0 1px 2px rgba(16,24,40,.04)}
.sect .s-tag{flex:none;font-size:11px;font-weight:700;color:#fff;background:var(--accent);
     border-radius:999px;padding:3px 10px;letter-spacing:.04em}
.sect .s-t{font-size:16px;font-weight:700;color:#152238;line-height:1.5;min-width:0}
.sect .s-m{margin-left:auto;flex:none;font-size:12px;color:var(--muted);
     font-variant-numeric:tabular-nums}
.sect input.np-edit-t{flex:1 1 auto;min-width:0;font:inherit;font-size:16px;font-weight:700;
     color:#152238;border:1px dashed #9dbcf5;border-radius:8px;padding:4px 8px;
     background:#fff;outline:0}
.notes .sect-in{margin:0 0 10px;padding:5px 9px;border-left:3px solid var(--accent);
     background:#f3f7ff;border-radius:6px;font-size:12.5px;color:#2b4b86;line-height:1.6;
     font-family:"Microsoft YaHei",system-ui,sans-serif}
/* ---- slide cards --------------------------------------------------------- */
.slide{background:var(--card);border:1px solid var(--line);border-radius:14px;
     margin:0 0 26px;overflow:hidden;scroll-margin-top:66px;
     box-shadow:0 1px 2px rgba(16,24,40,.04),0 10px 28px -20px rgba(16,24,40,.3);
     transition:box-shadow .2s,border-color .2s}
.slide.active{border-color:var(--accent);box-shadow:0 0 0 3px rgba(37,99,235,.12),0 12px 30px -18px rgba(16,24,40,.35)}
.slide.mod{border-color:var(--amber)}
.slide .bar{display:flex;align-items:center;gap:10px;padding:9px 16px;background:var(--head);border-bottom:1px solid var(--line)}
.slide .bar .no{font-weight:700;font-size:13px}
.slide .bar .tm{margin-left:auto;color:var(--muted);font-size:12px;font-variant-numeric:tabular-nums}
.mod-dot{display:none;font-size:11px;line-height:1.6;padding:1px 8px;border-radius:999px;
     background:#fef3c7;color:#8a5a09;border:1px solid #f0d493;cursor:pointer}
body.editing .slide.mod .mod-dot{display:inline-block}
.body{display:flex;gap:18px;align-items:stretch;padding:16px 18px}
.stage{flex:1 1 58%;min-width:0;aspect-ratio:16/9;background:var(--stage);
     border-radius:9px;display:flex;align-items:center;justify-content:center;overflow:hidden}
.stage img{max-width:100%;max-height:100%;display:block;object-fit:contain}
.notes-wrap{flex:1 1 42%;min-width:0;min-height:0;position:relative}
.notes{position:absolute;inset:0;overflow-y:auto;padding:0 12px 0 2px;
     scrollbar-width:thin;scrollbar-color:var(--thumb) var(--track)}
.notes h4{margin:0 0 8px;font-size:12px;letter-spacing:.08em;color:var(--accent);
     font-weight:700;position:sticky;top:0;background:var(--card);padding:3px 0 6px;z-index:2}
.notes p{margin:0 0 10px;font-family:"Songti SC","SimSun",serif;font-size:15px;
     line-height:1.9;color:#20242b;text-align:justify}
.notes p.empty{color:var(--muted);font-family:inherit;font-style:italic;font-size:14px}
.notes textarea{display:block;width:100%;height:100%;border:0;outline:0;resize:none;
     background:transparent;color:#20242b;padding:0;
     font-family:"Songti SC","SimSun",serif;font-size:15px;line-height:1.9;
     scrollbar-width:thin;scrollbar-color:var(--thumb) var(--track)}
.notes::-webkit-scrollbar{width:10px}
.notes::-webkit-scrollbar-track{background:var(--track);border-radius:8px}
.notes::-webkit-scrollbar-thumb{background:var(--thumb);border-radius:8px;border:2px solid var(--track)}
.notes::-webkit-scrollbar-thumb:hover{background:#9aa6b6}
.notes textarea::-webkit-scrollbar{width:10px}
.notes textarea::-webkit-scrollbar-track{background:var(--track);border-radius:8px}
.notes textarea::-webkit-scrollbar-thumb{background:var(--thumb);border-radius:8px;border:2px solid var(--track)}
.toast{position:fixed;left:50%;bottom:26px;transform:translateX(-50%);z-index:99;
     background:#111827;color:#fff;font-size:13px;padding:9px 16px;border-radius:9px;
     opacity:0;pointer-events:none;transition:opacity .25s}
.toast.on{opacity:.95}
@media(max-width:1180px){.rail{flex-basis:150px}}
@media(max-width:960px){
  .layout{padding:14px 10px 60px}
  .rail{display:none}
  header.top .t{max-width:52vw}
}
@media(max-width:820px){
  .body{flex-direction:column}
  .notes-wrap{height:300px;flex:none}
  .sect{flex-wrap:wrap}
  .sect .s-m{margin-left:0;width:100%}
}
</style></head>
<body>
<header class="top">
  <button id="railToggle" title="显示 / 隐藏左侧页面总览">☰ 收起目录</button>
  <div class="t">__TITLE__</div>
  <span class="sp"></span>
  <button id="prev">上一页</button>
  <button id="next">下一页</button>
  <span class="jump">
    <input id="jumpN" type="number" min="1" max="__COUNT__" placeholder="页码" inputmode="numeric">
    <span class="of">/ __COUNT__</span>
    <button id="jumpGo">跳转</button>
  </span>
  <button id="editToggle">编辑</button>
  <div class="c"><span id="cur">1</span> / __COUNT__</div>
  <div class="tools">
    <button id="btnSave" class="pri">保存（覆盖原文件）</button>
    <button id="btnSaveAs">另存为…</button>
    <button id="btnRevertAll">还原全部</button>
    <button id="btnCopy">复制全部讲稿</button>
    <span class="stat" id="stat"></span>
    <span class="badge" id="envBadge"></span>
  </div>
</header>
<div class="restore" id="restore">
  <span id="restoreMsg"></span>
  <button id="btnDropDraft">丢弃草稿</button>
</div>
<div class="layout">
  <aside class="rail" id="rail" aria-label="页面总览">__RAIL__</aside>
  <main class="wrap" id="deck">
__CARDS__
  </main>
</div>
<div class="toast" id="toast"></div>
<script>
/* ---- 导航 / 左目录 / 页码跳转 -------------------------------------------- */
var SECTS=__SECTIONS__;
var slides=Array.prototype.slice.call(document.querySelectorAll('.slide'));
var thumbs=Array.prototype.slice.call(document.querySelectorAll('.thumb'));
var idx=0, cur=document.getElementById('cur'), rail=document.getElementById('rail');
function setActive(i){
  if(i<0||i>=slides.length){return;}
  slides.forEach(function(s){s.classList.remove('active');});
  slides[i].classList.add('active');
  idx=i;cur.textContent=i+1;
  var th=thumbs[i];
  if(th){
    thumbs.forEach(function(t){t.classList.remove('active');});
    th.classList.add('active');
    if(rail){
      var rt=rail.getBoundingClientRect(), tt=th.getBoundingClientRect();
      if(tt.top<rt.top){rail.scrollTop+=tt.top-rt.top-6;}
      else if(tt.bottom>rt.bottom){rail.scrollTop+=tt.bottom-rt.bottom+6;}
    }
  }
}
function go(i){
  i=Math.max(0,Math.min(slides.length-1,i));
  slides[i].scrollIntoView({behavior:'smooth',block:'start'});
  setActive(i);
}
document.getElementById('prev').onclick=function(){go(idx-1);};
document.getElementById('next').onclick=function(){go(idx+1);};
window.addEventListener('keydown',function(e){
  var t=e.target, tag=t&&t.tagName;
  if(tag==='INPUT'||tag==='TEXTAREA'||(t&&t.isContentEditable)){return;}
  if(e.key==='ArrowRight'||e.key==='PageDown'){go(idx+1);}
  else if(e.key==='ArrowLeft'||e.key==='PageUp'){go(idx-1);}});
var io=new IntersectionObserver(function(es){es.forEach(function(e){
  if(e.isIntersecting){setActive(slides.indexOf(e.target));}});},{threshold:.5});
slides.forEach(function(s){io.observe(s);});
if(rail){
  rail.addEventListener('click',function(e){
    var el=e.target;
    var j=el.closest?el.closest('[data-jump]'):null;
    if(!j){return;}
    e.preventDefault();
    for(var i=0;i<slides.length;i++){if(slides[i].getAttribute('data-n')===j.getAttribute('data-jump')){
      go(i);return;}}
  });
}
function setRail(on,persist){
  document.body.classList.toggle('rail-hidden',!on);
  document.getElementById('railToggle').textContent=on?'☰ 收起目录':'☰ 目录';
  if(persist){try{window.localStorage.setItem(RAILKEY,on?'1':'0');}catch(err){}}
}
var RAILKEY='npedit::__DECKID__::rail';
var railInit=true;
try{if(window.localStorage.getItem(RAILKEY)==='0'){railInit=false;}}catch(err){}
setRail(railInit,false);
document.getElementById('railToggle').addEventListener('click',function(){
  setRail(document.body.classList.contains('rail-hidden'),true);});
document.getElementById('jumpGo').addEventListener('click',function(){
  var box=document.getElementById('jumpN');
  var v=parseInt(box.value,10);
  if(!v||v<1||v>slides.length){
    if(typeof window.__toast==='function'){window.__toast('请输入 1–'+slides.length+' 之间的页码');}
    box.focus();box.select();return;}
  go(v-1);box.blur();
});
document.getElementById('jumpN').addEventListener('keydown',function(e){
  if(e.key==='Enter'){e.preventDefault();document.getElementById('jumpGo').click();}});
</script>
<script>
/* ---- V2.2.0 编辑：讲稿 + 节标题，可「保存 / 另存为」 ----------------------
   原理：浏览器不允许网页静默写回本地文件，所以
     ·「另存为」调 showSaveFilePicker 选路径后写盘（Chrome/Edge 顶层页面）；
     ·「保存」用本次会话记住的文件句柄，一键覆盖，无需再选；
     · 环境不支持（非 Chromium 或在 iframe 预览窗口内）时自动退化为下载新文件；
     · 编辑过程中每次改动都会暂存草稿，误关页面可恢复。
   注意：textarea / input 的输入内容只存在于 DOM 的 value 属性里、不写回 HTML
        源码，因此导出前必须先 exitEdit() 把内容落成 p / span，再序列化。 */
(function(){
var DECK_FILE="__DECKFILE__", DECK_ID="__DECKID__";
var LSKEY="npedit::"+DECK_ID;
var CAN_FSA=(typeof window.showSaveFilePicker==="function");
var IN_FRAME=false;
try{IN_FRAME=(window.self!==window.top);}catch(err){IN_FRAME=true;}
var deck=document.getElementById("deck");
var secs=Array.prototype.slice.call(deck.querySelectorAll(".slide"));
var sectEls=Array.prototype.slice.call(deck.querySelectorAll(".sect[data-title]"));
var ORIG={}, ORIGT={}, editing=false, handle=null, dtimer=null;
secs.forEach(function(s,i){ORIG[i]=s.getAttribute("data-notes")||"";});
sectEls.forEach(function(e,i){ORIGT[i]=e.getAttribute("data-title")||"";});
var HINT={};
SECTS.forEach(function(s){
  if(s.hint_page&&s.hint_off!==null&&s.hint_off!==undefined&&s.hint_off>0){
    HINT[s.hint_page]={off:s.hint_off,idx:s.idx,label:s.label};}});

function esc(t){return String(t).replace(/[&<>]/g,function(c){
  return c==="&"?"&amp;":(c==="<"?"&lt;":"&gt;");});}
function paras(t){var o=[];String(t).split(/\n\s*\n/).forEach(function(b){
  b=b.split("\n").map(function(x){return x.trim();}).filter(function(x){return x;}).join(" ");
  if(b){o.push(b);}});return o;}
function psHTML(t){return paras(t).map(function(p){return "<p>"+esc(p)+"</p>";}).join("");}
function notesHTML(t,sec){
  var v=String(t==null?"":t), n=sec?(+sec.getAttribute("data-n")):0, h=HINT[n];
  if(h){
    var off=Math.max(0,Math.min(h.off,v.length));
    var L=psHTML(v.slice(0,off)), R=psHTML(v.slice(off));
    if(!L&&!R){return '<h4>演讲稿</h4><p class="empty">（本页无讲稿）</p>';}
    return '<h4>演讲稿</h4>'+L+
      '<div class="sect-in" data-sect="'+h.idx+'">'+esc(h.label)+'</div>'+R;
  }
  var body=psHTML(v);
  if(!body){return '<h4>演讲稿</h4><p class="empty">（本页无讲稿）</p>';}
  return '<h4>演讲稿</h4>'+body;
}
function curText(s){var ta=s.querySelector(".np-edit");return ta?ta.value:(s.getAttribute("data-notes")||"");}
function curTitle(e){var ip=e.querySelector(".np-edit-t");return ip?ip.value:(e.getAttribute("data-title")||"");}
function toast(m){var t=document.getElementById("toast");t.textContent=m;t.classList.add("on");
  clearTimeout(t._h);t._h=setTimeout(function(){t.classList.remove("on");},2200);}
window.__toast=toast;
function recount(){
  var c=0;
  secs.forEach(function(s,i){var d=curText(s)!==(ORIG[i]||"");s.classList.toggle("mod",d);if(d){c++;}});
  var tc=0;
  sectEls.forEach(function(e,i){if(curTitle(e)!==(ORIGT[i]||"")){tc++;}});
  document.getElementById("stat").textContent="已改动 "+c+" / "+secs.length+" 页"+
    (tc?("，节标题 "+tc+" 处"):"");
  return c+tc;
}
function isDirty(){
  for(var i=0;i<secs.length;i++){if(curText(secs[i])!==(ORIG[i]||"")){return true;}}
  for(var j=0;j<sectEls.length;j++){if(curTitle(sectEls[j])!==(ORIGT[j]||"")){return true;}}
  return false;
}
function syncEnv(){
  var b=document.getElementById("envBadge");
  document.getElementById("editToggle").textContent=editing?"退出编辑":"编辑";
  if(!CAN_FSA){b.className="badge warn";b.textContent="本环境不支持直接写盘，保存将改为下载新文件";}
  else if(IN_FRAME){b.className="badge warn";b.textContent="预览窗口内无法写盘，请在浏览器新标签页打开本文件";}
  else{b.className="badge ok";b.textContent="支持「保存」直接覆盖原文件";}
}
function enterEdit(){
  if(editing){return;}
  editing=true;document.body.classList.add("editing");
  secs.forEach(function(sec){
    var box=sec.querySelector(".notes");
    if(!box||box.querySelector(".np-edit")){return;}
    var ta=document.createElement("textarea");
    ta.className="np-edit";ta.spellcheck=false;
    ta.value=sec.getAttribute("data-notes")||"";
    ta.addEventListener("input",function(){recount();draftSoon();});
    box.innerHTML="";box.appendChild(ta);
  });
  sectEls.forEach(function(e){
    if(e.querySelector(".np-edit-t")){return;}
    var t=e.querySelector(".s-t");
    if(!t){return;}
    var ip=document.createElement("input");
    ip.className="np-edit-t";ip.type="text";ip.spellcheck=false;
    ip.value=e.getAttribute("data-title")||"";
    ip.addEventListener("input",function(){recount();draftSoon();});
    t.parentNode.replaceChild(ip,t);
  });
  recount();syncEnv();
}
function exitEdit(){
  if(!editing){return;}
  secs.forEach(function(sec){
    var ta=sec.querySelector(".np-edit");
    if(!ta){return;}
    var v=ta.value;sec.setAttribute("data-notes",v);
    sec.querySelector(".notes").innerHTML=notesHTML(v,sec);
  });
  sectEls.forEach(function(e){
    var ip=e.querySelector(".np-edit-t");
    if(!ip){return;}
    var v=ip.value;e.setAttribute("data-title",v);
    var sp=document.createElement("span");sp.className="s-t";sp.textContent=v;
    ip.parentNode.replaceChild(sp,ip);
  });
  editing=false;document.body.classList.remove("editing");syncEnv();
}
function buildOutput(){
  var was=editing;
  exitEdit();
  var c=document.documentElement.cloneNode(true);
  Array.prototype.slice.call(c.querySelectorAll(".slide.active")).forEach(function(s){s.classList.remove("active");});
  Array.prototype.slice.call(c.querySelectorAll(".thumb.active")).forEach(function(s){s.classList.remove("active");});
  var cu=c.querySelector("#cur");if(cu){cu.textContent="1";}
  var tt=c.querySelector("#toast");if(tt){tt.classList.remove("on");tt.textContent="";}
  var rs=c.querySelector("#restore");if(rs){rs.classList.remove("on");}
  var jn=c.querySelector("#jumpN");if(jn){jn.setAttribute("value","");}
  var out="<!DOCTYPE html>\n"+c.outerHTML;
  if(was){enterEdit();}
  return out;
}
function blobOf(html){return new Blob([html],{type:"text/html;charset=utf-8"});}
function download(){
  var a=document.createElement("a");
  a.href=URL.createObjectURL(blobOf(buildOutput()));
  a.download=DECK_FILE||"slides_viewer.html";
  document.body.appendChild(a);a.click();
  setTimeout(function(){URL.revokeObjectURL(a.href);a.remove();},6000);
  toast("已下载新文件（当前环境无法直接覆盖原文件）");
}
function markSaved(){
  secs.forEach(function(s,i){ORIG[i]=curText(s);s.classList.remove("mod");});
  sectEls.forEach(function(e,i){ORIGT[i]=curTitle(e);});
  try{window.localStorage.removeItem(LSKEY);}catch(err){}
  document.getElementById("restore").classList.remove("on");
  recount();
}
async function writeHandle(h){
  var w=await h.createWritable();
  await w.write(blobOf(buildOutput()));
  await w.close();
}
async function doSaveAs(){
  if(!CAN_FSA||IN_FRAME){download();return;}
  try{
    var h=await window.showSaveFilePicker({suggestedName:DECK_FILE||"slides_viewer.html",
      types:[{description:"HTML 文件",accept:{"text/html":[".html"]}}]});
    handle=h;await writeHandle(h);markSaved();toast("已保存：" + h.name);
  }catch(err){
    if(err&&err.name==="AbortError"){return;}
    toast("保存失败："+((err&&err.message)||err));
  }
}
async function doSave(){
  if(!CAN_FSA||IN_FRAME){download();return;}
  if(!handle){toast("首次保存请先选择文件位置");doSaveAs();return;}
  try{
    var st=await handle.requestPermission({mode:"readwrite"});
    if(st!=="granted"){doSaveAs();return;}
    await writeHandle(handle);markSaved();toast("已覆盖原文件：" + handle.name);
  }catch(err){
    if(err&&err.name==="AbortError"){return;}
    toast("保存失败："+((err&&err.message)||err));
  }
}
function draftSoon(){if(!editing){return;}clearTimeout(dtimer);dtimer=setTimeout(persistDraft,600);}
function persistDraft(){
  try{
    var o={},c=0;
    secs.forEach(function(s,i){var t=curText(s);if(t!==(ORIG[i]||"")){o[i]=t;c++;}});
    var tt={},tc=0;
    sectEls.forEach(function(e,i){var v=curTitle(e);if(v!==(ORIGT[i]||"")){tt["T"+i]=v;tc++;}});
    if(c||tc){window.localStorage.setItem(LSKEY,JSON.stringify({notes:o,titles:tt}));}
    else{window.localStorage.removeItem(LSKEY);}
  }catch(err){}
}
function loadDraft(){
  try{
    var raw=window.localStorage.getItem(LSKEY);
    if(!raw){return;}
    var p=JSON.parse(raw);
    var notes=(p&&p.notes)?p.notes:(p||{}), titles=(p&&p.titles)||{};
    var kn=Object.keys(notes), kt=Object.keys(titles);
    if(!kn.length&&!kt.length){return;}
    kn.forEach(function(i){
      var sec=secs[+i];if(!sec){return;}
      sec.setAttribute("data-notes",notes[i]);
      sec.querySelector(".notes").innerHTML=notesHTML(notes[i],sec);
    });
    kt.forEach(function(k){
      var e=sectEls[+String(k).slice(1)];if(!e){return;}
      e.setAttribute("data-title",titles[k]);
      var t=e.querySelector(".s-t");if(t){t.textContent=titles[k];}
    });
    document.getElementById("restoreMsg").textContent=
      "检测到上次未保存的草稿（"+kn.length+" 页讲稿 / "+kt.length+" 处节标题），已自动恢复。点「保存」写回文件，或丢弃草稿。";
    document.getElementById("restore").classList.add("on");
  }catch(err){}
}
function applyOrig(sec,i){
  var t=ORIG[i]||"";sec.setAttribute("data-notes",t);
  var ta=sec.querySelector(".np-edit");
  if(ta){ta.value=t;}else{sec.querySelector(".notes").innerHTML=notesHTML(t,sec);}
}
function applyOrigTitles(){
  sectEls.forEach(function(e,i){
    var v=ORIGT[i]||"";e.setAttribute("data-title",v);
    var ip=e.querySelector(".np-edit-t");
    if(ip){ip.value=v;return;}
    var t=e.querySelector(".s-t");
    if(t){t.textContent=v;return;}
    var sp=document.createElement("span");sp.className="s-t";sp.textContent=v;e.appendChild(sp);
  });
}
function dropDraft(banner){
  secs.forEach(applyOrig);applyOrigTitles();
  try{window.localStorage.removeItem(LSKEY);}catch(err){}
  if(banner){document.getElementById("restore").classList.remove("on");}
  recount();
}
document.getElementById("editToggle").addEventListener("click",function(){
  if(editing){exitEdit();}else{enterEdit();}});
document.getElementById("btnSave").addEventListener("click",doSave);
document.getElementById("btnSaveAs").addEventListener("click",doSaveAs);
document.getElementById("btnRevertAll").addEventListener("click",function(){
  if(!window.confirm("把全部讲稿与节标题还原为最初内容？（已保存到文件的修改不受影响）")){return;}
  dropDraft(true);toast("已还原全部内容");
});
document.getElementById("btnDropDraft").addEventListener("click",function(){
  dropDraft(true);toast("草稿已丢弃");
});
document.getElementById("btnCopy").addEventListener("click",function(){
  var byFrom={};
  SECTS.forEach(function(s){byFrom[s.from_n]=s;});
  var parts=[];
  secs.forEach(function(s,i){
    var n=+s.getAttribute("data-n");
    if(byFrom[n]){parts.push("# 【新增节 "+byFrom[n].idx+"】"+byFrom[n].title);}
    parts.push("## 第 "+n+" 页\n\n"+curText(s));
  });
  var txt=parts.join("\n\n\n");
  function ok(){toast("已复制 " + secs.length + " 页讲稿（含节标题）");}
  function fb(){
    var t=document.createElement("textarea");
    t.value=txt;t.style.position="fixed";t.style.left="-9999px";
    document.body.appendChild(t);t.select();
    try{document.execCommand("copy");ok();}catch(err){toast("复制失败，请手动选择文本");}
    t.remove();
  }
  if(navigator.clipboard&&navigator.clipboard.writeText){
    navigator.clipboard.writeText(txt).then(ok,fb);
  }else{fb();}
});
deck.addEventListener("click",function(e){
  var d=(e.target&&e.target.closest)?e.target.closest(".mod-dot"):null;
  if(!d){return;}
  var sec=d.closest(".slide"),i=secs.indexOf(sec);
  if(i<0){return;}
  if(!window.confirm("将本页讲稿还原为最初内容？")){return;}
  applyOrig(sec,i);recount();draftSoon();
});
window.addEventListener("beforeunload",function(e){
  if(isDirty()){e.preventDefault();e.returnValue="";}
});
loadDraft();syncEnv();recount();
})();
</script>
</body></html>
"""


def build_html(meta, out_path, title="幻灯片与讲稿", embed_dir="slides",
               max_width=1500, jpeg_quality=80, inline_images=True,
               sections=None, rail=True, rail_max_width=260, rail_quality=58,
               name_fixes=None):
    """Render every slide as a PPT-style stage on the LEFT with the 演讲稿
    (speaker notes) on the RIGHT, so slide and script read side by side.

    V2.2.0 additions
      * sections — pass output/sections.json (or its parsed dict). Each section
        gets a 【新增节】 divider before its first slide, a heading in the
        left thumbnail rail, and (when the host's hand-over sentence can be
        located inside a page) an inline 「下一节」 marker between paragraphs.
        Section titles are editable in the page's editor.
      * left thumbnail rail — dedicated small JPEGs (default 260px q58) inlined
        as data URIs; ~+7% file size. Grouped by section, click to jump, follows
        the scroll position, and can be collapsed (remembered per deck).
      * page-number jump box in the header, and 编辑 / 退出编辑 wording.
    """
    out_dir = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(out_dir, exist_ok=True)
    secs = load_sections(sections)
    by_from = {s["from_n"]: s for s in secs}
    anchors = locate_anchors(meta, secs, fixes=name_fixes) if secs else {}
    # A hint is only meaningful when the hand-over was spoken on an EARLIER page
    # than the section's first slide — otherwise the marker would sit inside the
    # section it announces.
    hints = {}
    for sec in secs:
        a = anchors.get(sec["idx"])
        if a and a["page"] < sec["from_n"]:
            hints[a["page"]] = {
                "off": a["off"],
                "html": '<div class="sect-in" data-sect="%d">下一节 · %s</div>'
                        % (sec["idx"], html.escape(sec["title"])),
            }

    embed_abs = None
    if not inline_images and embed_dir:
        embed_abs = os.path.join(out_dir, embed_dir)
        os.makedirs(embed_abs, exist_ok=True)

    def ref_for(src, n, prefix, mw, q, lazy=""):
        if not src or not os.path.exists(src):
            return ""
        if inline_images:
            try:
                b = _jpeg_bytes(src, mw, q)
                return "data:image/jpeg;base64," + base64.b64encode(b).decode("ascii")
            except Exception:
                pass
        if embed_abs:
            dst_name = "%s%03d.jpg" % (prefix, n)
            try:
                with open(os.path.join(embed_abs, dst_name), "wb") as f:
                    f.write(_jpeg_bytes(src, mw, q))
                return "%s/%s" % (embed_dir, dst_name)
            except Exception:
                pass
        return os.path.relpath(os.path.abspath(src), out_dir).replace("\\", "/")

    def notes_inner(raw, hint):
        """Paragraph markup. When a mid-page hint exists the raw text is split at
        the hand-over's end so the 「下一节」 line lands BETWEEN two paragraphs —
        the same split the in-page editor reproduces from the stored offset."""
        if hint and hint.get("off") is not None:
            off = max(0, min(int(hint["off"]), len(raw)))
            a = _notes_paragraphs(raw[:off])
            b = _notes_paragraphs(raw[off:])
            if not a and not b:
                return '<p class="empty">（本页无讲稿）</p>'
            return ("".join("<p>%s</p>" % html.escape(p) for p in a)
                    + hint["html"]
                    + "".join("<p>%s</p>" % html.escape(p) for p in b))
        paras = _notes_paragraphs(raw)
        if not paras:
            return '<p class="empty">（本页无讲稿）</p>'
        return "".join("<p>%s</p>" % html.escape(p) for p in paras)

    cards, rail_items = [], []
    for i, s in enumerate(meta, 1):
        n = s.get("n", i)
        sec = by_from.get(n)
        if sec:
            cards.append(
                '<section class="sect" id="sec%d" data-sect="%d" data-title="%s">'
                '<span class="s-tag">新增节</span>'
                '<span class="s-t">%s</span>'
                '<span class="s-m">%s</span></section>' % (
                    sec["idx"], sec["idx"], html.escape(sec["title"], quote=True),
                    html.escape(sec["title"]), html.escape(section_range_label(meta, sec))))
            rail_items.append(
                '<div class="rail-sect" data-jump="%d" title="跳到本节第 1 页">'
                '<span class="n">节 %d</span>%s</div>' % (
                    n, sec["idx"], html.escape(sec["title"])))

        main_ref = ref_for(s.get("file", ""), n, "slide_", max_width, jpeg_quality)
        thumb_ref = ref_for(s.get("file", ""), n, "thumb_", rail_max_width, rail_quality) \
            if rail else ""
        ts = fmt_mmss(s.get("t"))
        raw_notes = _norm_notes(s.get("notes", ""))

        hint = hints.get(n)
        notes_html = notes_inner(raw_notes, hint)
        lazy = "" if inline_images else ' loading="lazy"'
        img_html = ('<img src="%s" alt="第 %d 页"%s>' % (html.escape(main_ref), n, lazy)) \
            if main_ref else '<em style="color:#889">（无图片）</em>'
        cards.append(
            '<section class="slide" id="s%d" data-n="%d" data-notes="%s">'
            '<div class="bar"><span class="no">第 %d 页</span>'
            '<span class="mod-dot" title="本页已改动，点此还原">已改</span>'
            '<span class="tm">%s</span></div>'
            '<div class="body">'
            '<div class="stage">%s</div>'
            '<div class="notes-wrap"><div class="notes"><h4>演讲稿</h4>%s</div></div>'
            '</div></section>' % (
                n, n, html.escape(raw_notes, quote=True), n, ts, img_html, notes_html))
        if rail and thumb_ref:
            rail_items.append(
                '<a class="thumb" href="#s%d" data-jump="%d">'
                '<img src="%s" alt="第 %d 页" loading="lazy"><span class="tn">%d</span></a>'
                % (n, n, html.escape(thumb_ref), n, n))

    # payload for the in-page editor: rail hint placement + section titles
    payload = []
    for sec in secs:
        a = anchors.get(sec["idx"], {})
        usable = bool(a) and a["page"] < sec["from_n"]
        payload.append({
            "idx": sec["idx"], "title": sec["title"], "from_n": sec["from_n"],
            "hint_page": a.get("page") if usable else None,
            "hint_off": a.get("off") if usable else None,
            "label": "下一节 · %s" % sec["title"],
        })

    deck_file = os.path.basename(os.path.abspath(out_path))
    deck_id = hashlib.md5(("%s|%d|%s|v%s" % (title, len(meta), deck_file, DECK_VERSION))
                          .encode("utf-8")).hexdigest()[:12]
    page = (_HTML_TMPL.replace("__TITLE__", html.escape(title))
            .replace("__COUNT__", str(len(meta)))
            .replace("__DECKFILE__", html.escape(deck_file, quote=True))
            .replace("__DECKID__", deck_id)
            .replace("__SECTIONS__", json.dumps(payload, ensure_ascii=True))
            .replace("__RAIL__", "\n".join(rail_items) if rail else "")
            .replace("__CARDS__", "\n".join(cards)))
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(page)
    return out_path


def _jpeg_bytes(src, max_width, quality):
    from PIL import Image
    import io as _io
    im = Image.open(src).convert("RGB")
    if max_width and im.width > max_width:
        im = im.resize((max_width, int(im.height * max_width / im.width)), Image.LANCZOS)
    buf = _io.BytesIO()
    im.save(buf, "JPEG", quality=quality, optimize=True)
    return buf.getvalue()


if __name__ == "__main__":
    # Convenience: rebuild deliverables from an existing meta (no re-alignment).
    ap = argparse.ArgumentParser()
    ap.add_argument("--meta", required=True)
    ap.add_argument("--ocr", default="output/slides_ocr.txt")
    ap.add_argument("--outdir", default="output")
    ap.add_argument("--title", default="幻灯片与讲稿")
    ap.add_argument("--sections", default=None,
                    help="output/sections.json (from build_sections.py)")
    ap.add_argument("--separate-images", action="store_true",
                    help="write slide images to <outdir>/slides/ instead of inlining "
                         "(default: inline as ONE self-contained HTML, safe to share)")
    ap.add_argument("--no-rail", action="store_true",
                    help="do not build the left thumbnail rail")
    ap.add_argument("--notes-pptx", action="store_true",
                    help="[legacy] also build slides_with_notes.pptx (slides + speaker "
                         "notes). NOT part of the V2.x deliverable set — the transcript "
                         "belongs in the HTML viewer; only use this if someone "
                         "explicitly asks for a notes deck")
    ap.add_argument("--max-width", type=int, default=1500)
    ap.add_argument("--quality", type=int, default=80)
    args = ap.parse_args()
    meta = load_meta(args.meta)
    ocr_by_slide = parse_ocr(args.ocr)
    secs = load_sections(args.sections)
    fixes = load_name_fixes(args.sections)
    if fixes:
        print("讲者姓名校正 %s -> 应用 %d 处" % (fixes, apply_name_fixes(meta, fixes)))
    os.makedirs(args.outdir, exist_ok=True)
    md = build_markdown(meta, ocr_by_slide,
                        os.path.join(args.outdir, "transcript_from_slides.md"),
                        sections=secs, title=args.title)
    tr = build_transcript_docx(meta, os.path.join(args.outdir, "transcript_by_speaker.docx"),
                               sections=secs, title="逐字稿 按讲者分节")
    cp_ = build_clean_pptx(meta, os.path.join(args.outdir, "slides.pptx"), sections=secs)
    vw = build_html(meta, os.path.join(args.outdir, "slides_viewer.html"), title=args.title,
                    max_width=args.max_width, jpeg_quality=args.quality,
                    inline_images=not args.separate_images,
                    sections=secs, rail=not args.no_rail, name_fixes=fixes)
    print("逐字稿     -> %s\n自查对照   -> %s\n幻灯片 PPT -> %s\nViewer HTML -> %s"
          % (tr, md, cp_, vw))
    if args.notes_pptx:
        p = build_notes_pptx(meta, os.path.join(args.outdir, "slides_with_notes.pptx"),
                             sections=secs)
        print("[legacy] Notes PPTX -> %s" % p)
