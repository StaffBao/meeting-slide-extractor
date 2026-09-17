# build_html_from_pptx.py  (V2.1)
# Turn an EXISTING deck that already carries slide images + speaker notes
# (e.g. slides_with_notes.pptx produced by this skill) into the HTML viewer:
# slide on the LEFT (PPT-style "stage"), 演讲稿 on the RIGHT with its own
# vertical scroll track -- without re-running the whole video pipeline.
#
# Which images does the viewer need?
#   The HTML *displays* the slide image, so the images are inherently required.
#   Two ways to supply them:
#     * default             -> base64-embedded, producing ONE self-contained
#                              .html (no slides/ folder on disk) that survives
#                              being shared/copied on its own
#     * --separate-images   -> <outdir>/slides/slide_NNN.jpg, referenced by the
#                              HTML (smaller HTML, but the folder must travel
#                              with the file)
#
# Usage:
#   python build_html_from_pptx.py --pptx slides_with_notes.pptx \
#       --outdir output --title "幻灯片与讲稿" [--separate-images] \
#       [--max-width 1500] [--quality 80]
import argparse, io, os, shutil, tempfile
from pptx import Presentation
from PIL import Image
from build_notes_ppt import build_html


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--outdir", default="output")
    ap.add_argument("--out", default=None, help="output html (default <outdir>/slides_viewer.html)")
    ap.add_argument("--title", default="幻灯片与讲稿")
    ap.add_argument("--separate-images", action="store_true",
                    help="write slide images to <outdir>/slides/ instead of inlining "
                         "(default: inline as ONE self-contained HTML, safe to share)")
    ap.add_argument("--max-width", type=int, default=1500)
    ap.add_argument("--quality", type=int, default=80)
    args = ap.parse_args()

    inline = not args.separate_images

    # where to stage the downscaled JPEGs
    if inline:
        src_dir = tempfile.mkdtemp(prefix="viewer_src_")
    else:
        src_dir = os.path.join(args.outdir, "slides")
        os.makedirs(src_dir, exist_ok=True)

    prs = Presentation(args.pptx)
    meta = []
    for i, slide in enumerate(prs.slides, 1):
        blob = None
        for shp in slide.shapes:
            try:
                b = shp.image.blob
            except Exception:
                continue
            if blob is None or len(b) > len(blob):
                blob = b
        rel = ""
        if blob is not None:
            im = Image.open(io.BytesIO(blob)).convert("RGB")
            if im.width > args.max_width:
                im = im.resize((args.max_width, int(im.height * args.max_width / im.width)),
                               Image.LANCZOS)
            dst = os.path.join(src_dir, f"slide_{i:03d}.jpg")
            im.save(dst, "JPEG", quality=args.quality, optimize=True)
            rel = dst
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text
        except Exception:
            pass
        meta.append({"n": i, "file": rel, "t": None, "notes": notes})

    out = args.out or os.path.join(args.outdir, "slides_viewer.html")
    vw = build_html(meta, out, title=args.title,
                    embed_dir=None, inline_images=inline,
                    max_width=args.max_width, jpeg_quality=args.quality)
    if inline:
        shutil.rmtree(src_dir, ignore_errors=True)
    n_files = len(meta)
    extra = "  (self-contained single file)" if inline else "  (+ slides/ folder)"
    print(f"slides: {n_files}  ->  {vw}{extra}")


if __name__ == "__main__":
    main()
